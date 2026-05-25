import csv
import json
import re
import uuid
from io import BytesIO, StringIO
from pathlib import Path

from datetime import datetime

from fastapi import APIRouter, Cookie, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.database import get_db
from app.models import (
    AuditLog,
    Competitor,
    CompetitorVehicleAssignment,
    Conflict,
    Criterion,
    CriteriaSet,
    CurrentEventState,
    Delay,
    Event,
    EventSettings,
    GraphicsState,
    GraphicsAction,
    Judge,
    JudgeIssueReport,
    JudgeSession,
    Pad,
    PinAccount,
    RecoveryImport,
    Run,
    RunTimer,
    ScoreChangeRequest,
    ScoreEntry,
    ScoreItemEntry,
    Vehicle,
    now_utc,
)
from app.schemas import (
    AdminStatusActionInput,
    CompetitorImportCommit,
    CompetitorInput,
    CriterionInput,
    ConnectivitySettingsInput,
    DelayInput,
    EventLifecycleInput,
    EventSettingsInput,
    GraphicsStateInput,
    PinAccountInput,
    RecoveryImportInput,
    RunOrderInput,
    ScoreOverrideInput,
    ScoreVoidInput,
    TimerControlInput,
)
from app.seed import seed_initial_data
from app.services.audit import log_action
from app.services.backups import backup_health as backup_health_snapshot
from app.services.backups import build_backup_payload, write_backup
from app.services.results import RESULT_MODE_LABELS, competitor_scoreboard, normalize_result_mode
from app.services.scoring import calculate_points
from app.routers.public import DEFAULT_CONNECTIVITY, connectivity_payload

router = APIRouter(prefix="/api/admin", tags=["admin"])

LOGO_UPLOAD_DIR = Path("app/static/uploads/logos")
LOGO_URL_PREFIX = "/static/uploads/logos"
ALLOWED_LOGO_TYPES = {"image/png"}
MAX_LOGO_BYTES = 5 * 1024 * 1024


BACKUP_MODELS = [
    Event,
    Pad,
    EventSettings,
    CriteriaSet,
    Criterion,
    PinAccount,
    Judge,
    JudgeSession,
    Competitor,
    Vehicle,
    Run,
    CompetitorVehicleAssignment,
    ScoreEntry,
    ScoreItemEntry,
    ScoreChangeRequest,
    Delay,
    CurrentEventState,
    JudgeIssueReport,
    GraphicsState,
    GraphicsAction,
    RunTimer,
    RecoveryImport,
    Conflict,
    AuditLog,
]

BACKUP_TABLES = {model.__tablename__: model for model in BACKUP_MODELS}

DEFAULT_CRITERIA_ROWS = [
    ("Instant Smoke", "Smoke", "numeric_score", 0, 10, 1, 1, 10),
    ("Smoke Volume", "Volume", "numeric_score", 0, 10, 1, 1, 20),
    ("Consistent Smoke", "Consistent", "numeric_score", 0, 10, 1, 1, 30),
    ("Pad Use", "Pad", "numeric_score", 0, 30, 1, 1, 40),
    ("Driver Skill", "Skill", "numeric_score", 0, 20, 1, 1, 50),
    ("Wall Taps", "Taps", "count_penalty", 0, 20, 1, 10, 60),
    ("Tyre Burst", "Tyres", "count_bonus", 0, 8, 1, 5, 70),
    ("Reverse", "Reverse", "toggle", 0, 1, 1, -5, 80),
    ("Drive off unassisted", "Drive Off", "toggle", 0, 1, 1, 10, 90),
]

DEFAULT_PIN_ACCOUNTS = [
    ("Judge 1", "101", "judge", 1),
    ("Judge 2", "102", "judge", 1),
    ("Graphics", "0201", "graphics", 2),
    ("Admin", "1234", "admin", 2),
    ("High Admin", "12345", "high_admin", 3),
    ("Owner", "123456", "owner", 4),
]

RUN_TYPES = {"competition", "demo", "fun", "exhibition", "test"}
INACTIVE_RUN_STATES = {"finished", "skipped", "withdrawn", "disqualified"}
DEFAULT_DELAY_PRESETS = [
    {"key": "lunch_break", "label": "Lunch Break", "message": "Lunch break", "minutes": 30},
    {"key": "track_cleanup", "label": "Track Cleanup", "message": "Track cleanup", "minutes": 5},
    {"key": "breakdown_delay", "label": "Breakdown Delay", "message": "Breakdown recovery", "minutes": 10},
]


def get_session(db: Session, session_id: int | None) -> JudgeSession | None:
    if session_id is None:
        return None
    session = db.get(JudgeSession, session_id)
    if not session or not session.active:
        raise HTTPException(status_code=401, detail="Invalid session")
    session.last_seen_at = now_utc()
    return session


def require_admin(session: JudgeSession | None) -> None:
    if not session or session.role not in {"admin", "high_admin", "owner"}:
        raise HTTPException(status_code=403, detail="Admin session required")


def require_graphics_or_admin(session: JudgeSession | None) -> None:
    if not session or session.role not in {"graphics", "high_admin", "owner"}:
        raise HTTPException(status_code=403, detail="Graphics, high-admin, or owner session required")


def require_owner(session: JudgeSession | None) -> None:
    if not session or session.role != "owner":
        raise HTTPException(status_code=403, detail="Owner session required")


def require_high_admin(session: JudgeSession | None) -> None:
    if not session or session.role not in {"high_admin", "owner"}:
        raise HTTPException(status_code=403, detail="High-admin or owner session required")


def expected_pin_length(role: str) -> int:
    return {"judge": 3, "graphics": 4, "admin": 4, "high_admin": 5, "owner": 6}[role]


def validate_pin_payload(payload: PinAccountInput) -> None:
    if payload.role not in {"judge", "graphics", "admin", "high_admin", "owner"}:
        raise HTTPException(status_code=400, detail="Invalid role")
    pin = (payload.pin or "").strip()
    if not pin.isdigit():
        raise HTTPException(status_code=400, detail="PIN must contain numbers only")
    required_length = expected_pin_length(payload.role)
    if len(pin) != required_length:
        raise HTTPException(
            status_code=400,
            detail=f"{payload.role.replace('_', ' ').title()} PIN must be {required_length} digits",
        )


def get_event_pad(db: Session, event_id: int) -> Pad:
    pad = db.scalar(select(Pad).where(Pad.event_id == event_id, Pad.active.is_(True)).order_by(Pad.id).limit(1))
    if pad:
        return pad
    pad = Pad(event_id=event_id, name="Pad 1", active=True)
    db.add(pad)
    db.flush()
    return pad


def active_event(db: Session) -> Event:
    event = db.scalar(select(Event).where(Event.active.is_(True)).order_by(Event.id.desc()).limit(1))
    if not event:
        raise HTTPException(status_code=404, detail="Active event not found")
    return event


def event_settings_record(db: Session, event_id: int) -> EventSettings:
    settings = db.scalar(select(EventSettings).where(EventSettings.event_id == event_id).limit(1))
    if settings:
        return settings
    settings = EventSettings(event_id=event_id)
    db.add(settings)
    db.flush()
    return settings


def clean_base_url(value: str | None) -> str:
    cleaned = (value or "").strip().rstrip("/")
    if not cleaned:
        return ""
    if not re.match(r"^https?://", cleaned, flags=re.IGNORECASE):
        raise HTTPException(status_code=400, detail="Connection URLs must start with http:// or https://")
    return cleaned


def graphics_state_record(db: Session, event_id: int) -> GraphicsState:
    state = db.scalar(select(GraphicsState).where(GraphicsState.event_id == event_id).limit(1))
    if state:
        return state
    pad = get_event_pad(db, event_id)
    state = GraphicsState(
        event_id=event_id,
        pad_id=pad.id,
        preset="standard_run",
        enabled_layers=["current_competitor", "judge_likeness"],
        layout_config={},
    )
    db.add(state)
    db.flush()
    return state


def clone_json(value):
    return json.loads(json.dumps(value)) if value is not None else {}


def normalize_delay_presets(value) -> list[dict]:
    source = value if isinstance(value, list) and value else DEFAULT_DELAY_PRESETS
    presets = []
    for idx, item in enumerate(source[:3]):
        raw = item if isinstance(item, dict) else {}
        fallback = DEFAULT_DELAY_PRESETS[idx] if idx < len(DEFAULT_DELAY_PRESETS) else DEFAULT_DELAY_PRESETS[-1]
        label = str(raw.get("label") or raw.get("name") or fallback["label"]).strip()[:40] or fallback["label"]
        message = str(raw.get("message") or label).strip()[:120] or label
        try:
            minutes = int(float(raw.get("minutes", raw.get("estimated_minutes", fallback["minutes"])) or 0))
        except (TypeError, ValueError):
            minutes = fallback["minutes"]
        key = str(raw.get("key") or re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_") or f"preset_{idx + 1}")
        presets.append(
            {
                "key": key[:40],
                "label": label,
                "message": message,
                "minutes": max(0, min(240, minutes)),
            }
        )
    return presets or clone_json(DEFAULT_DELAY_PRESETS)


def create_default_criteria_set(db: Session, event_id: int) -> CriteriaSet:
    criteria_set = CriteriaSet(event_id=event_id, name="Default Burnout Criteria", active=True)
    db.add(criteria_set)
    db.flush()
    for name, short, type_, min_v, max_v, step, points, order in DEFAULT_CRITERIA_ROWS:
        db.add(
            Criterion(
                criteria_set_id=criteria_set.id,
                name=name,
                short_label=short,
                type=type_,
                min_value=min_v,
                max_value=max_v,
                step_size=step,
                points_per_unit=points,
                required=False,
                display_order=order,
                active=True,
            )
        )
    return criteria_set


def create_default_pin_accounts(db: Session, event_id: int) -> None:
    for display_name, pin, role, level in DEFAULT_PIN_ACCOUNTS:
        account = PinAccount(
            event_id=event_id,
            display_name=display_name,
            pin=pin,
            role=role,
            permission_level=level,
            active=True,
        )
        db.add(account)
        db.flush()
        if role == "judge":
            db.add(Judge(event_id=event_id, pin_account_id=account.id, name=display_name))


def finish_other_active_runs(db: Session, run: Run) -> list[int]:
    active_states = {"current", "staging", "on_pad", "live"}
    others = db.scalars(
        select(Run).where(
            Run.event_id == run.event_id,
            Run.pad_id == run.pad_id,
            Run.id != run.id,
            Run.state.in_(active_states),
        )
    ).all()
    changed = []
    for other in others:
        other.state = "finished"
        changed.append(other.id)
    return changed


def next_run_after(db: Session, run: Run) -> Run | None:
    queue_position = run.queue_position or 0
    return db.scalar(
        select(Run)
        .where(
            Run.event_id == run.event_id,
            Run.pad_id == run.pad_id,
            Run.id != run.id,
            or_(
                Run.state.is_(None),
                Run.state.not_in(list(INACTIVE_RUN_STATES)),
            ),
            or_(
                Run.queue_position > queue_position,
                Run.queue_position.is_(None),
            ),
        )
        .order_by(Run.queue_position.is_(None), Run.queue_position, Run.id)
        .limit(1)
    )


def place_run_next(db: Session, session: JudgeSession, run: Run) -> dict:
    state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == session.event_id))
    current_run = db.get(Run, state.official_current_run_id) if state and state.official_current_run_id else None
    active_runs = db.scalars(
        select(Run)
        .where(
            Run.event_id == session.event_id,
            Run.pad_id == run.pad_id,
            or_(Run.state.is_(None), Run.state.not_in(list(INACTIVE_RUN_STATES))),
        )
        .order_by(Run.queue_position.is_(None), Run.queue_position, Run.id)
    ).all()
    if run not in active_runs:
        active_runs.append(run)
    ordered = [item for item in active_runs if item.id != run.id]
    if current_run:
        current_index = next((idx for idx, item in enumerate(ordered) if item.id == current_run.id), -1)
        insert_at = current_index + 1 if current_index >= 0 else 0
    else:
        insert_at = 0
    ordered.insert(insert_at, run)
    changed_positions = []
    for idx, item in enumerate(ordered, start=1):
        if item.queue_position != idx:
            changed_positions.append({"run_id": item.id, "old": item.queue_position, "new": idx})
            item.queue_position = idx
    if run.state in INACTIVE_RUN_STATES or not run.state:
        run.state = "queued"
    return {
        "current_run_id": current_run.id if current_run else None,
        "next_run_id": run.id,
        "changed_positions": changed_positions,
    }


def active_timer(db: Session, event_id: int, pad_id: int) -> RunTimer | None:
    return db.scalar(
        select(RunTimer)
        .where(
            RunTimer.event_id == event_id,
            RunTimer.pad_id == pad_id,
            RunTimer.status.in_(["ready", "running", "stopped"]),
        )
        .order_by(RunTimer.id.desc())
        .limit(1)
    )


def as_float(value) -> float | None:
    return float(value) if value is not None else None


def format_competitor(run: Run | None) -> str:
    if not run or not run.competitor:
        return "No competitor"
    return f"#{run.competitor.entry_number} {run.competitor.driver_name}"


def run_payload(run: Run | None) -> dict | None:
    if not run:
        return None
    return {
        "id": run.id,
        "queue_position": run.queue_position,
        "run_number": run.run_number,
        "heat_number": run.heat_number,
        "run_type": run.run_type,
        "state": run.state,
        "include_in_results": run.include_in_results,
        "competitor": {
            "id": run.competitor.id,
            "entry_number": run.competitor.entry_number,
            "driver_name": run.competitor.driver_name,
            "class_name": run.competitor.class_name,
            "status": run.competitor.status,
        },
        "vehicle": {
            "id": run.vehicle.id if run.vehicle else None,
            "name": run.vehicle.name if run.vehicle else "",
            "engine": run.vehicle.engine if run.vehicle else "",
            "plate": run.vehicle.plate if run.vehicle else "",
        },
    }


def score_item_payload(item: ScoreItemEntry) -> dict:
    criterion = item.criterion
    return {
        "id": item.id,
        "criterion_id": item.criterion_id,
        "criterion_name": criterion.name if criterion else f"Criterion {item.criterion_id}",
        "short_label": criterion.short_label if criterion else None,
        "display_order": criterion.display_order if criterion else 9999,
        "type": criterion.type if criterion else None,
        "value": as_float(item.value),
        "calculated_points": as_float(item.calculated_points),
        "intentionally_blank": bool(item.intentionally_blank),
    }


def entry_context_maps(db: Session, event_id: int, entries: list[ScoreEntry]) -> tuple[dict[int, Run], dict[int, Judge]]:
    run_ids = sorted({entry.run_id for entry in entries if entry.run_id})
    judge_ids = sorted({entry.judge_id for entry in entries if entry.judge_id})
    runs_by_id = {}
    judges_by_id = {}
    if run_ids:
        runs_by_id = {
            run.id: run
            for run in db.scalars(select(Run).where(Run.event_id == event_id, Run.id.in_(run_ids))).all()
        }
    if judge_ids:
        judges_by_id = {
            judge.id: judge
            for judge in db.scalars(select(Judge).where(Judge.event_id == event_id, Judge.id.in_(judge_ids))).all()
        }
    return runs_by_id, judges_by_id


def score_entry_payload(entry: ScoreEntry, runs_by_id: dict[int, Run], judges_by_id: dict[int, Judge]) -> dict:
    run = runs_by_id.get(entry.run_id)
    judge = judges_by_id.get(entry.judge_id) if entry.judge_id else None
    items = sorted(
        [score_item_payload(item) for item in entry.items],
        key=lambda item: (item["display_order"], item["criterion_id"]),
    )
    return {
        "id": entry.id,
        "run_id": entry.run_id,
        "run": run_payload(run),
        "heat_number": run.heat_number if run else None,
        "competitor": format_competitor(run),
        "vehicle": run.vehicle.name if run and run.vehicle else "",
        "judge_id": entry.judge_id,
        "judge_name": judge.name if judge else (f"Judge {entry.judge_id}" if entry.judge_id else "Manual/Admin"),
        "device_id": entry.device_id,
        "status": entry.status,
        "source_type": entry.source_type,
        "original_score_entry_id": entry.original_score_entry_id,
        "active_for_results": bool(entry.active_for_results),
        "reason": entry.reason,
        "total": as_float(entry.total),
        "updated_at": entry.updated_at.isoformat(),
        "submitted_at": entry.submitted_at.isoformat() if entry.submitted_at else None,
        "items": items,
        "blank_items": sum(1 for item in items if item["value"] is None or item["intentionally_blank"]),
    }


def issue_report_payload(db: Session, report: JudgeIssueReport) -> dict:
    run = db.get(Run, report.run_id) if report.run_id else None
    judge = db.get(Judge, report.judge_id)
    return {
        "id": report.id,
        "status": report.status,
        "issue_type": report.issue_type,
        "note": report.note,
        "device_id": report.device_id,
        "judge_id": report.judge_id,
        "judge_name": judge.name if judge else f"Judge {report.judge_id}",
        "run_id": report.run_id,
        "run": run_payload(run),
        "competitor": format_competitor(run),
        "created_at": report.created_at.isoformat(),
        "updated_at": report.updated_at.isoformat(),
    }


def score_change_request_payload(db: Session, request: ScoreChangeRequest) -> dict:
    run = db.get(Run, request.run_id)
    current_at_request = db.get(Run, request.current_run_id_at_request) if request.current_run_id_at_request else None
    judge = db.get(Judge, request.judge_id)
    original = None
    if request.original_score_entry_id:
        entry = db.get(ScoreEntry, request.original_score_entry_id)
        if entry:
            runs_by_id, judges_by_id = entry_context_maps(db, request.event_id, [entry])
            original = score_entry_payload(entry, runs_by_id, judges_by_id)
    return {
        "id": request.id,
        "status": request.status,
        "reason": request.reason,
        "requested_changes": request.requested_changes,
        "device_id": request.device_id,
        "judge_id": request.judge_id,
        "judge_name": judge.name if judge else f"Judge {request.judge_id}",
        "run_id": request.run_id,
        "run": run_payload(run),
        "competitor": format_competitor(run),
        "current_run_at_request": run_payload(current_at_request),
        "original_score_entry_id": request.original_score_entry_id,
        "original_score": original,
        "created_at": request.created_at.isoformat(),
        "updated_at": request.updated_at.isoformat(),
        "resolved_at": request.resolved_at.isoformat() if request.resolved_at else None,
    }


def criteria_totals_payload(entries: list[ScoreEntry]) -> list[dict]:
    totals: dict[int, dict] = {}
    for entry in entries:
        for item in entry.items:
            criterion = item.criterion
            bucket = totals.setdefault(
                item.criterion_id,
                {
                    "criterion_id": item.criterion_id,
                    "criterion_name": criterion.name if criterion else f"Criterion {item.criterion_id}",
                    "short_label": criterion.short_label if criterion else None,
                    "display_order": criterion.display_order if criterion else 9999,
                    "total_points": 0.0,
                    "entered_count": 0,
                    "blank_count": 0,
                },
            )
            if item.value is None or item.intentionally_blank:
                bucket["blank_count"] += 1
            else:
                bucket["entered_count"] += 1
            if item.calculated_points is not None:
                bucket["total_points"] += float(item.calculated_points)
    return sorted(totals.values(), key=lambda item: (item["display_order"], item["criterion_id"]))


def timer_elapsed_seconds(timer: RunTimer) -> int:
    elapsed = int(timer.elapsed_offset_seconds or 0)
    if timer.status == "running" and timer.started_at:
        elapsed += max(0, int((now_utc() - timer.started_at).total_seconds()))
    return elapsed


def start_run_timer(db: Session, session: JudgeSession, run: Run) -> RunTimer:
    timer = active_timer(db, run.event_id, run.pad_id)
    if timer and timer.run_id != run.id and timer.status == "running":
        timer.elapsed_offset_seconds = timer_elapsed_seconds(timer)
        timer.stopped_at = now_utc()
        timer.status = "stopped"
    if timer and timer.run_id is None:
        timer.run_id = run.id
    elif not timer or timer.run_id != run.id:
        mode = timer.mode if timer else "count_up"
        target_seconds = timer.target_seconds if timer else 90
        timer = RunTimer(
            event_id=run.event_id,
            pad_id=run.pad_id,
            run_id=run.id,
            mode=mode,
            target_seconds=target_seconds,
            status="ready",
            source="run_control",
        )
        db.add(timer)
        db.flush()
    if timer.status != "running":
        timer.started_at = now_utc()
        timer.stopped_at = None
        timer.status = "running"
        log_action(
            db,
            event_id=session.event_id,
            actor_pin_account_id=session.pin_account_id,
            device_id=session.device_id,
            action_type="timer_started",
            entity_type="run_timer",
            entity_id=timer.id,
            details={"run_id": run.id, "mode": timer.mode, "target_seconds": timer.target_seconds},
        )
    return timer


def stop_run_timer(db: Session, session: JudgeSession, run: Run) -> RunTimer | None:
    timer = active_timer(db, run.event_id, run.pad_id)
    if not timer or timer.run_id != run.id:
        return None
    if timer.status == "running":
        timer.elapsed_offset_seconds = timer_elapsed_seconds(timer)
        timer.stopped_at = now_utc()
        timer.status = "stopped"
        log_action(
            db,
            event_id=session.event_id,
            actor_pin_account_id=session.pin_account_id,
            device_id=session.device_id,
            action_type="timer_stopped",
            entity_type="run_timer",
            entity_id=timer.id,
            details={"run_id": run.id, "elapsed_seconds": timer.elapsed_offset_seconds},
        )
    return timer


@router.post("/current-run/{run_id}")
def set_current_run(run_id: int, session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_admin(session)

    run = db.get(Run, run_id)
    if not run:
        raise HTTPException(status_code=404, detail="Run not found")

    state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == session.event_id))
    old_run_id = state.official_current_run_id if state else None
    if state:
        state.official_current_run_id = run.id
        state.pad_id = run.pad_id

    if old_run_id and old_run_id != run.id:
        old_run = db.get(Run, old_run_id)
        if old_run and old_run.state == "current":
            old_run.state = None

    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="current_run_changed",
        entity_type="run",
        entity_id=run.id,
        details={"old_run_id": old_run_id, "new_run_id": run.id},
    )
    db.commit()
    return {"ok": True, "current_run_id": run.id}


@router.post("/next-run/{run_id}")
def set_next_run(run_id: int, session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_admin(session)
    run = db.get(Run, run_id)
    if not run or run.event_id != session.event_id:
        raise HTTPException(status_code=404, detail="Run not found")
    result = place_run_next(db, session, run)
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="next_competitor_set",
        entity_type="run",
        entity_id=run.id,
        details=result,
    )
    db.commit()
    return {"ok": True, **result}


@router.post("/current-run/clear")
def clear_current_run(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_admin(session)

    state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == session.event_id))
    old_run_id = state.official_current_run_id if state else None
    if state:
        state.official_current_run_id = None

    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="current_run_cleared",
        entity_type="current_event_state",
        details={"old_run_id": old_run_id},
    )
    db.commit()
    return {"ok": True, "old_run_id": old_run_id}


@router.post("/current-state/clear-current-run")
def clear_current_run_state(session_id: int, db: Session = Depends(get_db)) -> dict:
    return clear_current_run(session_id, db)


@router.post("/run/{run_id}/state/{state_name}")
def set_run_state(
    run_id: int,
    state_name: str,
    session_id: int,
    start_timer: bool = True,
    db: Session = Depends(get_db),
) -> dict:
    session = get_session(db, session_id)
    require_admin(session)

    if state_name == "none":
        state_name = ""
    allowed = {
        "",
        "queued",
        "staging",
        "current",
        "on_pad",
        "live",
        "finished",
        "skipped",
        "delayed",
        "withdrawn",
        "mechanical_issue",
        "disqualified",
    }
    state_name = state_name or ""
    if state_name not in allowed:
        raise HTTPException(status_code=400, detail="Invalid run state")
    run = db.get(Run, run_id)
    if not run:
        raise HTTPException(status_code=404, detail="Run not found")
    old = run.state
    run.state = state_name
    auto_finished = []
    advanced_to_run_id = None
    if state_name in {"staging", "on_pad", "live"}:
        auto_finished = finish_other_active_runs(db, run)
    if state_name == "live" or (state_name == "on_pad" and start_timer):
        start_run_timer(db, session, run)
    if state_name == "finished":
        stop_run_timer(db, session, run)
        event_state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == session.event_id))
        next_run = next_run_after(db, run)
        if event_state and next_run and event_state.official_current_run_id == run.id:
            event_state.official_current_run_id = next_run.id
            event_state.pad_id = next_run.pad_id
            advanced_to_run_id = next_run.id
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="run_state_changed",
        entity_type="run",
        entity_id=run.id,
        details={
            "old_state": old,
            "new_state": state_name,
            "auto_finished_run_ids": auto_finished,
            "advanced_to_run_id": advanced_to_run_id,
            "start_timer_requested": start_timer,
        },
    )
    db.commit()
    return {"ok": True, "run_id": run.id, "state": run.state, "advanced_to_run_id": advanced_to_run_id}


@router.post("/timer")
def control_timer(payload: TimerControlInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_admin(session)

    event_id = session.event_id
    pad = get_event_pad(db, event_id)
    run = db.get(Run, payload.run_id) if payload.run_id else None
    if not run:
        state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == event_id))
        run = db.get(Run, state.official_current_run_id) if state and state.official_current_run_id else None

    timer = active_timer(db, event_id, pad.id)
    if not timer or (run and timer.run_id != run.id and payload.action in {"set", "start", "reset"}):
        timer = RunTimer(event_id=event_id, pad_id=pad.id, run_id=run.id if run else None)
        db.add(timer)
        db.flush()

    old_elapsed = timer_elapsed_seconds(timer)
    timer.mode = payload.mode if payload.mode in {"count_up", "count_down"} else "count_up"
    timer.target_seconds = max(0, int(payload.target_seconds or 0))

    if payload.action == "start":
        if payload.elapsed_seconds is not None:
            timer.elapsed_offset_seconds = max(0, int(payload.elapsed_seconds))
        timer.started_at = now_utc()
        timer.stopped_at = None
        timer.status = "running"
    elif payload.action == "stop":
        timer.elapsed_offset_seconds = old_elapsed
        timer.stopped_at = now_utc()
        timer.status = "stopped"
    elif payload.action == "reset":
        timer.elapsed_offset_seconds = 0
        timer.started_at = None
        timer.stopped_at = None
        timer.status = "ready"
    elif payload.action == "adjust":
        timer.elapsed_offset_seconds = max(0, int(payload.elapsed_seconds or 0))
        if timer.status == "running":
            timer.started_at = now_utc()
    else:
        if payload.elapsed_seconds is not None:
            timer.elapsed_offset_seconds = max(0, int(payload.elapsed_seconds))
            if timer.status == "running":
                timer.started_at = now_utc()

    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type=f"timer_{payload.action}",
        entity_type="run_timer",
        entity_id=timer.id,
        details={
            "run_id": timer.run_id,
            "mode": timer.mode,
            "target_seconds": timer.target_seconds,
            "elapsed_seconds": timer_elapsed_seconds(timer),
        },
    )
    db.commit()
    return {"ok": True, "timer_id": timer.id}


@router.post("/delay")
def create_delay(payload: DelayInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    if payload.official:
        require_admin(session)
    else:
        require_graphics_or_admin(session)

    event_id = session.event_id if session else 1
    pad = get_event_pad(db, event_id)
    delay = Delay(
        event_id=event_id,
        pad_id=pad.id,
        created_by_pin_account_id=session.pin_account_id if session else None,
        created_by_role=session.role if session else None,
        source=payload.source,
        message=payload.message,
        estimated_minutes=payload.estimated_minutes,
        estimated_seconds=payload.estimated_seconds,
        visibility=payload.visibility,
        official=payload.official,
        priority=payload.priority,
    )
    db.add(delay)
    db.flush()
    state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == event_id))
    if state:
        if payload.official:
            state.official_delay_id = delay.id
        else:
            state.graphics_delay_id = delay.id
    log_action(
        db,
        event_id=event_id,
        actor_pin_account_id=session.pin_account_id if session else None,
        device_id=session.device_id if session else None,
        action_type="delay_created",
        entity_type="delay",
        entity_id=delay.id,
        details=payload.model_dump(),
    )
    db.commit()
    return {"ok": True, "delay_id": delay.id}


@router.post("/delay/clear")
def clear_delays(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_admin(session)
    delays = db.scalars(select(Delay).where(Delay.event_id == session.event_id, Delay.active.is_(True))).all()
    for delay in delays:
        delay.active = False
        delay.resolved_at = now_utc()
    state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == session.event_id))
    if state:
        state.official_delay_id = None
        state.graphics_delay_id = None
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="delay_cleared",
        entity_type="delay",
        details={"count": len(delays)},
    )
    db.commit()
    return {"ok": True, "cleared": len(delays)}


@router.post("/graphics-state")
def update_graphics_state(
    payload: GraphicsStateInput,
    kairix_session_id: str | None = Cookie(default=None),
    db: Session = Depends(get_db),
) -> dict:
    session_id = payload.session_id
    if session_id is None and kairix_session_id and kairix_session_id.isdigit():
        session_id = int(kairix_session_id)
    session = get_session(db, session_id)
    require_graphics_or_admin(session)
    state = db.scalar(select(GraphicsState).where(GraphicsState.event_id == session.event_id))
    if not state:
        raise HTTPException(status_code=404, detail="Graphics state not found")
    state.preset = payload.preset
    state.enabled_layers = payload.enabled_layers
    state.manual_message = payload.manual_message
    state.fun_detail_mode = payload.fun_detail_mode
    state.show_debug_zones = payload.show_debug_zones
    if payload.layout_config is not None:
        state.layout_config = payload.layout_config
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="graphics_state_updated",
        entity_type="graphics_state",
        entity_id=state.id,
        details=payload.model_dump(),
    )
    db.commit()
    return {"ok": True}


@router.get("/settings")
def get_admin_settings(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_high_admin(session)
    event = db.get(Event, session.event_id) or active_event(db)
    settings = event_settings_record(db, event.id)
    graphics = graphics_state_record(db, event.id)
    return {
        "event": {
            "id": event.id,
            "name": event.name,
            "venue": event.venue,
            "event_date": event.event_date,
            "active": event.active,
        },
        "settings": {
            "result_mode": normalize_result_mode(settings.result_mode),
            "result_mode_label": RESULT_MODE_LABELS.get(normalize_result_mode(settings.result_mode), "All heats"),
            "public_delay_seconds": settings.public_delay_seconds,
            "show_public_total_scores": settings.show_public_total_scores,
            "show_graphics_total_scores": settings.show_graphics_total_scores,
            "judge_likeness_enabled": settings.judge_likeness_enabled,
            "landing_notice": settings.landing_notice,
            "delay_presets": normalize_delay_presets(settings.delay_presets),
        },
        "theme": (graphics.layout_config or {}).get("theme") or {},
    }


@router.put("/settings")
def update_admin_settings(payload: EventSettingsInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_high_admin(session)
    event = db.get(Event, session.event_id) or active_event(db)
    settings = event_settings_record(db, event.id)
    old = {
        "name": event.name,
        "venue": event.venue,
        "event_date": event.event_date,
        "result_mode": settings.result_mode,
        "public_delay_seconds": settings.public_delay_seconds,
        "show_public_total_scores": settings.show_public_total_scores,
        "show_graphics_total_scores": settings.show_graphics_total_scores,
        "judge_likeness_enabled": settings.judge_likeness_enabled,
        "landing_notice": settings.landing_notice,
        "delay_presets": normalize_delay_presets(settings.delay_presets),
    }
    event.name = (payload.event_name or "").strip() or event.name
    event.venue = (payload.venue or "").strip() or None
    event.event_date = (payload.event_date or "").strip() or None
    settings.result_mode = normalize_result_mode(payload.result_mode)
    settings.public_delay_seconds = max(0, int(payload.public_delay_seconds or 0))
    settings.show_public_total_scores = bool(payload.show_public_total_scores)
    settings.show_graphics_total_scores = bool(payload.show_graphics_total_scores)
    settings.judge_likeness_enabled = bool(payload.judge_likeness_enabled)
    settings.landing_notice = (payload.landing_notice or "").strip() or None
    if session.role == "owner":
        settings.delay_presets = normalize_delay_presets(payload.delay_presets)
    log_action(
        db,
        event_id=event.id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="event_settings_updated",
        entity_type="event_settings",
        entity_id=settings.id,
        details={"old": old, "new": payload.model_dump(exclude={"session_id"})},
    )
    db.commit()
    return {"ok": True}


@router.post("/result-mode/{result_mode}")
def update_result_mode(result_mode: str, session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_high_admin(session)
    event = db.get(Event, session.event_id) or active_event(db)
    settings = event_settings_record(db, event.id)
    old = settings.result_mode
    settings.result_mode = normalize_result_mode(result_mode)
    log_action(
        db,
        event_id=event.id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="result_mode_changed",
        entity_type="event_settings",
        entity_id=settings.id,
        details={"old": old, "new": settings.result_mode},
    )
    db.commit()
    return {
        "ok": True,
        "result_mode": settings.result_mode,
        "result_mode_label": RESULT_MODE_LABELS[settings.result_mode],
    }


@router.get("/connectivity")
def get_connectivity_settings(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    event = active_event(db)
    settings = event_settings_record(db, event.id)
    return {"ok": True, "connectivity": connectivity_payload(settings)}


@router.put("/connectivity")
def update_connectivity_settings(payload: ConnectivitySettingsInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_owner(session)
    event = active_event(db)
    settings = event_settings_record(db, event.id)
    old = connectivity_payload(settings)
    config = dict(DEFAULT_CONNECTIVITY)
    config.update(
        {
            "local_base_url": clean_base_url(payload.local_base_url),
            "cloud_base_url": clean_base_url(payload.cloud_base_url),
            "fallback_enabled": bool(payload.fallback_enabled),
            "prefer_local": bool(payload.prefer_local),
            "auto_return_to_local": bool(payload.auto_return_to_local),
            "request_timeout_ms": max(500, min(15000, int(payload.request_timeout_ms or 2500))),
            "retry_local_after_seconds": max(5, min(600, int(payload.retry_local_after_seconds or 20))),
            "health_path": (payload.health_path or "/api/health").strip() or "/api/health",
        }
    )
    if not config["health_path"].startswith("/"):
        config["health_path"] = f"/{config['health_path']}"
    settings.connectivity_config = config
    log_action(
        db,
        event_id=event.id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="connectivity_settings_updated",
        entity_type="event_settings",
        entity_id=settings.id,
        details={"old": old, "new": config},
    )
    db.commit()
    return {"ok": True, "connectivity": config}


@router.get("/sessions/active")
def active_sessions_summary(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_high_admin(session)
    sessions = db.scalars(
        select(JudgeSession)
        .where(JudgeSession.event_id == session.event_id, JudgeSession.active.is_(True))
        .order_by(JudgeSession.last_seen_at.desc(), JudgeSession.id.desc())
    ).all()
    accounts = {
        account.id: account
        for account in db.scalars(select(PinAccount).where(PinAccount.event_id == session.event_id)).all()
    }
    judges = {
        judge.id: judge
        for judge in db.scalars(select(Judge).where(Judge.event_id == session.event_id)).all()
    }
    now = now_utc()
    rows = []
    connected = idle = stale = 0
    hidden_old = 0
    hide_after_seconds = 3 * 60 * 60
    for item in sessions:
        age_seconds = max(0, int((now - item.last_seen_at).total_seconds()))
        if age_seconds <= 45:
            status = "connected"
        elif age_seconds <= 300:
            status = "idle"
        else:
            status = "stale"
        if age_seconds > hide_after_seconds:
            hidden_old += 1
            continue
        if status == "connected":
            connected += 1
        elif status == "idle":
            idle += 1
        else:
            stale += 1
        account = accounts.get(item.pin_account_id)
        judge = judges.get(item.judge_id) if item.judge_id else None
        session_age_seconds = max(0, int((now - item.created_at).total_seconds())) if item.created_at else None
        rows.append(
            {
                "session_id": item.id,
                "role": item.role,
                "display_name": account.display_name if account else item.role.replace("_", " ").title(),
                "pin_account_id": item.pin_account_id,
                "account_role": account.role if account else item.role,
                "permission_level": account.permission_level if account else None,
                "judge_name": judge.name if judge else None,
                "device_id": item.device_id,
                "device_short_id": item.device_id[-10:] if item.device_id else None,
                "created_at": item.created_at.isoformat() if item.created_at else None,
                "last_seen_at": item.last_seen_at.isoformat(),
                "session_age_seconds": session_age_seconds,
                "seconds_since_seen": age_seconds,
                "status": status,
            }
        )
    return {
        "summary": {
            "total": len(rows),
            "connected": connected,
            "idle": idle,
            "stale": stale,
            "hidden_old": hidden_old,
            "hide_after_seconds": hide_after_seconds,
        },
        "sessions": rows,
    }


@router.get("/graphics/layout-export")
def export_graphics_layout(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_graphics_or_admin(session)
    state = graphics_state_record(db, session.event_id)
    event = db.get(Event, session.event_id)
    return {
        "schema": "kairix_graphics_layout_v1",
        "exported_at": now_utc().isoformat(),
        "event": {
            "id": event.id if event else session.event_id,
            "name": event.name if event else None,
            "venue": event.venue if event else None,
            "event_date": event.event_date if event else None,
        },
        "graphics": {
            "preset": state.preset,
            "enabled_layers": state.enabled_layers or [],
            "fun_detail_mode": state.fun_detail_mode,
            "show_debug_zones": state.show_debug_zones,
            "layout_config": state.layout_config or {},
        },
    }


@router.post("/graphics/layout-import")
async def import_graphics_layout(
    session_id: int = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> dict:
    session = get_session(db, session_id)
    require_graphics_or_admin(session)
    content = await file.read()
    try:
        payload = json.loads(content.decode("utf-8"))
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail="Graphics layout file must be valid JSON") from exc
    graphics_payload = payload.get("graphics")
    if not isinstance(graphics_payload, dict):
        raise HTTPException(status_code=400, detail="Graphics layout file is missing a graphics payload")

    state = graphics_state_record(db, session.event_id)
    state.preset = str(graphics_payload.get("preset") or state.preset or "layers")
    if isinstance(graphics_payload.get("enabled_layers"), list):
        state.enabled_layers = [str(item) for item in graphics_payload.get("enabled_layers") or []]
    state.fun_detail_mode = bool(graphics_payload.get("fun_detail_mode", state.fun_detail_mode))
    state.show_debug_zones = bool(graphics_payload.get("show_debug_zones", state.show_debug_zones))
    if isinstance(graphics_payload.get("layout_config"), dict):
        state.layout_config = graphics_payload["layout_config"]
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="graphics_layout_imported",
        entity_type="graphics_state",
        entity_id=state.id,
        details={"filename": file.filename, "schema": payload.get("schema")},
    )
    db.commit()
    return {"ok": True}


@router.post("/events/archive-and-create")
def archive_and_create_event(payload: EventLifecycleInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_owner(session)
    current_event = active_event(db)
    current_settings = event_settings_record(db, current_event.id)
    current_graphics = graphics_state_record(db, current_event.id)
    current_pad = get_event_pad(db, current_event.id)

    active_sessions = db.scalars(
        select(JudgeSession).where(JudgeSession.event_id == current_event.id, JudgeSession.active.is_(True))
    ).all()
    for active_session in active_sessions:
        active_session.active = False

    current_event.active = False
    new_event = Event(
        name=(payload.event_name or "").strip() or "Burnout Competition",
        venue=(payload.venue or "").strip() or current_event.venue,
        event_date=(payload.event_date or "").strip() or None,
        active=True,
    )
    db.add(new_event)
    db.flush()

    new_pad = Pad(event_id=new_event.id, name=current_pad.name if current_pad else "Pad 1", active=True)
    db.add(new_pad)
    db.flush()

    db.add(
        EventSettings(
            event_id=new_event.id,
            result_mode=current_settings.result_mode,
            public_delay_seconds=current_settings.public_delay_seconds,
            show_public_total_scores=current_settings.show_public_total_scores,
            show_graphics_total_scores=current_settings.show_graphics_total_scores,
            judge_likeness_enabled=current_settings.judge_likeness_enabled,
            allow_submit_with_nulls=current_settings.allow_submit_with_nulls,
            max_judges=current_settings.max_judges,
            landing_notice=current_settings.landing_notice if payload.copy_notice else None,
            connectivity_config=clone_json(current_settings.connectivity_config or {}),
            delay_presets=clone_json(current_settings.delay_presets or []),
        )
    )
    db.add(CurrentEventState(event_id=new_event.id, pad_id=new_pad.id, current_heat=1))

    copied_layout = clone_json(current_graphics.layout_config or {}) if payload.copy_graphics_layout else {}
    copied_layers = list(current_graphics.enabled_layers or []) if payload.copy_graphics_layout else [
        "current_competitor",
        "judge_likeness",
    ]
    db.add(
        GraphicsState(
            event_id=new_event.id,
            pad_id=new_pad.id,
            preset=current_graphics.preset if payload.copy_graphics_layout else "standard_run",
            enabled_layers=copied_layers,
            manual_message=None,
            fun_detail_mode=current_graphics.fun_detail_mode if payload.copy_graphics_layout else False,
            show_debug_zones=current_graphics.show_debug_zones if payload.copy_graphics_layout else False,
            layout_config=copied_layout,
        )
    )
    db.add(RunTimer(event_id=new_event.id, pad_id=new_pad.id, mode="count_up", target_seconds=90, status="ready"))

    if payload.copy_criteria:
        current_set = db.scalar(
            select(CriteriaSet)
            .where(CriteriaSet.event_id == current_event.id, CriteriaSet.active.is_(True))
            .order_by(CriteriaSet.id.desc())
            .limit(1)
        )
        if current_set:
            new_set = CriteriaSet(event_id=new_event.id, name=current_set.name, active=True)
            db.add(new_set)
            db.flush()
            current_items = db.scalars(
                select(Criterion)
                .where(Criterion.criteria_set_id == current_set.id, Criterion.active.is_(True))
                .order_by(Criterion.display_order, Criterion.id)
            ).all()
            for item in current_items:
                db.add(
                    Criterion(
                        criteria_set_id=new_set.id,
                        name=item.name,
                        short_label=item.short_label,
                        description=item.description,
                        type=item.type,
                        min_value=item.min_value,
                        max_value=item.max_value,
                        step_size=item.step_size,
                        points_per_unit=item.points_per_unit,
                        allows_negative=item.allows_negative,
                        required=item.required,
                        visible_to_judges=item.visible_to_judges,
                        visible_to_graphics=item.visible_to_graphics,
                        display_order=item.display_order,
                        active=item.active,
                    )
                )
        else:
            create_default_criteria_set(db, new_event.id)
    else:
        create_default_criteria_set(db, new_event.id)

    if payload.copy_pin_accounts:
        current_accounts = db.scalars(
            select(PinAccount)
            .where(PinAccount.event_id == current_event.id, PinAccount.active.is_(True))
            .order_by(PinAccount.permission_level, PinAccount.role, PinAccount.id)
        ).all()
        for account in current_accounts:
            new_account = PinAccount(
                event_id=new_event.id,
                display_name=account.display_name,
                pin=account.pin,
                role=account.role,
                permission_level=account.permission_level,
                active=account.active,
            )
            db.add(new_account)
            db.flush()
            if account.role == "judge":
                judge = db.scalar(
                    select(Judge).where(Judge.pin_account_id == account.id, Judge.active.is_(True)).limit(1)
                )
                db.add(
                    Judge(
                        event_id=new_event.id,
                        pin_account_id=new_account.id,
                        name=judge.name if judge else account.display_name,
                        status="active",
                        active=True,
                    )
                )
    else:
        create_default_pin_accounts(db, new_event.id)

    log_action(
        db,
        event_id=new_event.id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="event_archived_and_created",
        entity_type="event",
        entity_id=new_event.id,
        details={
            "archived_event_id": current_event.id,
            "new_event_name": new_event.name,
            "copy_criteria": payload.copy_criteria,
            "copy_pin_accounts": payload.copy_pin_accounts,
            "copy_graphics_layout": payload.copy_graphics_layout,
            "copy_notice": payload.copy_notice,
            "deactivated_sessions": len(active_sessions),
        },
    )
    db.commit()
    return {
        "ok": True,
        "new_event_id": new_event.id,
        "archived_event_id": current_event.id,
        "deactivated_sessions": len(active_sessions),
        "message": "New active event created. Log in again to use the new event.",
    }


def logo_file_payload(path: Path) -> dict:
    stat = path.stat()
    return {
        "filename": path.name,
        "url": f"{LOGO_URL_PREFIX}/{path.name}",
        "size": stat.st_size,
        "uploaded_at": datetime.fromtimestamp(stat.st_mtime).isoformat(),
    }


def safe_logo_path(filename: str) -> Path:
    name = Path(filename).name
    if not re.fullmatch(r"[A-Za-z0-9_.-]+\.png", name):
        raise HTTPException(status_code=400, detail="Logo filename is invalid")
    path = (LOGO_UPLOAD_DIR / name).resolve()
    root = LOGO_UPLOAD_DIR.resolve()
    if root not in path.parents:
        raise HTTPException(status_code=400, detail="Logo path is invalid")
    return path


@router.get("/graphics/logos")
def list_graphics_logos(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_graphics_or_admin(session)
    LOGO_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    files = sorted(
        (logo_file_payload(path) for path in LOGO_UPLOAD_DIR.glob("*.png") if path.is_file()),
        key=lambda item: item["uploaded_at"],
        reverse=True,
    )
    return {"logos": files}


@router.post("/graphics/logos")
async def upload_graphics_logo(
    session_id: int = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> dict:
    session = get_session(db, session_id)
    require_graphics_or_admin(session)
    if file.content_type not in ALLOWED_LOGO_TYPES:
        raise HTTPException(status_code=400, detail="Logo must be a PNG file")
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Logo file is empty")
    if len(content) > MAX_LOGO_BYTES:
        raise HTTPException(status_code=400, detail="Logo file must be 5 MB or smaller")
    if not content.startswith(b"\x89PNG\r\n\x1a\n"):
        raise HTTPException(status_code=400, detail="Logo must be a valid PNG file")

    LOGO_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    original = Path(file.filename or "logo.png").stem
    slug = re.sub(r"[^A-Za-z0-9_.-]+", "-", original).strip("-._") or "logo"
    filename = f"{slug[:48]}-{uuid.uuid4().hex[:8]}.png"
    path = safe_logo_path(filename)
    path.write_bytes(content)
    payload = logo_file_payload(path)
    state = db.scalar(select(GraphicsState).where(GraphicsState.event_id == session.event_id))
    if state:
        config = dict(state.layout_config or {})
        logo_style = dict(config.get("logo_style") or {})
        logo_style["url"] = payload["url"]
        config["logo_style"] = logo_style
        state.layout_config = config
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="graphics_logo_uploaded",
        entity_type="graphics_logo",
        details=payload,
    )
    db.commit()
    return {"ok": True, "logo": payload}


@router.delete("/graphics/logos/{filename}")
def delete_graphics_logo(filename: str, session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_graphics_or_admin(session)
    path = safe_logo_path(filename)
    if not path.exists():
        raise HTTPException(status_code=404, detail="Logo not found")
    path.unlink()
    state = db.scalar(select(GraphicsState).where(GraphicsState.event_id == session.event_id))
    url = f"{LOGO_URL_PREFIX}/{Path(filename).name}"
    if state and isinstance(state.layout_config, dict):
        config = dict(state.layout_config)
        logo_style = dict(config.get("logo_style") or {})
        if logo_style.get("url") == url:
            logo_style["url"] = ""
            config["logo_style"] = logo_style
            state.layout_config = config
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="graphics_logo_deleted",
        entity_type="graphics_logo",
        details={"filename": Path(filename).name, "url": url},
    )
    db.commit()
    return {"ok": True}


@router.post("/graphics-featured-run/{run_id}")
def set_graphics_featured_run(run_id: int, session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_graphics_or_admin(session)
    run = db.get(Run, run_id)
    if not run:
        raise HTTPException(status_code=404, detail="Run not found")
    state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == session.event_id))
    if state:
        state.graphics_featured_run_id = run.id
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="graphics_featured_run_changed",
        entity_type="run",
        entity_id=run.id,
        details={"run_id": run.id},
    )
    db.commit()
    return {"ok": True, "graphics_featured_run_id": run.id}


@router.post("/graphics-featured-run/clear")
def clear_graphics_featured_run(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_graphics_or_admin(session)

    state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == session.event_id))
    old_run_id = state.graphics_featured_run_id if state else None
    if state:
        state.graphics_featured_run_id = None

    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="graphics_featured_run_cleared",
        entity_type="current_event_state",
        details={"old_run_id": old_run_id},
    )
    db.commit()
    return {"ok": True, "old_run_id": old_run_id}


@router.post("/current-state/clear-graphics-featured-run")
def clear_graphics_featured_run_state(session_id: int, db: Session = Depends(get_db)) -> dict:
    return clear_graphics_featured_run(session_id, db)


@router.get("/live-scoring")
def live_scoring(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_admin(session)
    state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == session.event_id))
    if not state or not state.official_current_run_id:
        return {"current_run_id": None, "current_run": None, "scores": [], "recent_scores": []}
    current_run = db.get(Run, state.official_current_run_id)
    current_entries = db.scalars(
        select(ScoreEntry)
        .options(selectinload(ScoreEntry.items).selectinload(ScoreItemEntry.criterion))
        .where(
            ScoreEntry.run_id == state.official_current_run_id,
            ScoreEntry.active_for_results.is_(True),
        )
        .order_by(ScoreEntry.updated_at.desc())
    ).all()
    recent_entries = db.scalars(
        select(ScoreEntry)
        .options(selectinload(ScoreEntry.items).selectinload(ScoreItemEntry.criterion))
        .where(
            ScoreEntry.event_id == session.event_id,
            ScoreEntry.active_for_results.is_(True),
        )
        .order_by(ScoreEntry.updated_at.desc())
        .limit(20)
    ).all()
    runs_by_id, judges_by_id = entry_context_maps(db, session.event_id, current_entries + recent_entries)
    if current_run:
        runs_by_id[current_run.id] = current_run

    active_judges = db.scalars(
        select(Judge).where(Judge.event_id == session.event_id, Judge.active.is_(True)).order_by(Judge.id)
    ).all()
    current_by_judge = {}
    if state.official_current_run_id:
        for entry in current_entries:
            if entry.judge_id:
                existing = current_by_judge.get(entry.judge_id)
                if not existing or entry.updated_at > existing.updated_at:
                    current_by_judge[entry.judge_id] = entry
    latest_sessions = db.scalars(
        select(JudgeSession)
        .where(JudgeSession.event_id == session.event_id, JudgeSession.judge_id.is_not(None), JudgeSession.active.is_(True))
        .order_by(JudgeSession.last_seen_at.desc())
    ).all()
    session_by_judge = {}
    for judge_session in latest_sessions:
        session_by_judge.setdefault(judge_session.judge_id, judge_session)

    return {
        "current_run_id": state.official_current_run_id,
        "current_run": run_payload(current_run),
        "scores": [score_entry_payload(entry, runs_by_id, judges_by_id) for entry in current_entries],
        "recent_scores": [score_entry_payload(entry, runs_by_id, judges_by_id) for entry in recent_entries],
        "judge_statuses": [
            {
                "judge_id": judge.id,
                "judge_name": judge.name,
                "attendance_status": judge.status,
                "active": judge.active,
                "score_status": current_by_judge.get(judge.id).status if current_by_judge.get(judge.id) else "missing",
                "score_entry_id": current_by_judge.get(judge.id).id if current_by_judge.get(judge.id) else None,
                "last_seen_at": session_by_judge.get(judge.id).last_seen_at.isoformat()
                if session_by_judge.get(judge.id)
                else None,
            }
            for judge in active_judges
        ],
    }


@router.get("/judge-alerts")
def judge_alerts(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_admin(session)
    issues = db.scalars(
        select(JudgeIssueReport)
        .where(
            JudgeIssueReport.event_id == session.event_id,
            JudgeIssueReport.status.in_(["open", "pending"]),
        )
        .order_by(JudgeIssueReport.created_at.desc())
        .limit(25)
    ).all()
    changes = db.scalars(
        select(ScoreChangeRequest)
        .where(ScoreChangeRequest.event_id == session.event_id, ScoreChangeRequest.status == "pending")
        .order_by(ScoreChangeRequest.created_at.desc())
        .limit(25)
    ).all()
    return {
        "issues": [issue_report_payload(db, report) for report in issues],
        "score_change_requests": [score_change_request_payload(db, request) for request in changes],
        "counts": {"issues": len(issues), "score_change_requests": len(changes), "total": len(issues) + len(changes)},
    }


@router.post("/judge-issue-reports/{report_id}/status")
def update_issue_report_status(report_id: int, payload: AdminStatusActionInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_admin(session)
    if payload.status not in {"open", "acknowledged", "actioned", "dismissed"}:
        raise HTTPException(status_code=400, detail="Invalid issue report status")
    report = db.get(JudgeIssueReport, report_id)
    if not report or report.event_id != session.event_id:
        raise HTTPException(status_code=404, detail="Issue report not found")
    old_status = report.status
    report.status = payload.status
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="judge_issue_report_status_changed",
        entity_type="judge_issue_report",
        entity_id=report.id,
        details={"old": old_status, "new": report.status, "note": payload.note},
    )
    db.commit()
    return {"ok": True, "report": issue_report_payload(db, report)}


@router.post("/score-change-requests/{request_id}/status")
def update_score_change_request_status(
    request_id: int, payload: AdminStatusActionInput, db: Session = Depends(get_db)
) -> dict:
    session = get_session(db, payload.session_id)
    require_admin(session)
    if payload.status not in {"pending", "approved", "rejected", "resolved_manually"}:
        raise HTTPException(status_code=400, detail="Invalid score change request status")
    request = db.get(ScoreChangeRequest, request_id)
    if not request or request.event_id != session.event_id:
        raise HTTPException(status_code=404, detail="Score change request not found")
    old_status = request.status
    request.status = payload.status
    if payload.status != "pending":
        request.resolved_by_pin_account_id = session.pin_account_id
        request.resolved_at = now_utc()
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="score_change_request_status_changed",
        entity_type="score_change_request",
        entity_id=request.id,
        details={"old": old_status, "new": request.status, "note": payload.note},
    )
    db.commit()
    return {"ok": True, "request": score_change_request_payload(db, request)}


@router.post("/score-overrides")
def create_score_override(payload: ScoreOverrideInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_high_admin(session)
    original = db.get(ScoreEntry, payload.original_score_entry_id)
    if not original or original.event_id != session.event_id:
        raise HTTPException(status_code=404, detail="Original score entry not found")
    run = db.get(Run, original.run_id)
    if not run:
        raise HTTPException(status_code=404, detail="Run not found")
    source_type = "owner_override" if session.role == "owner" else "high_admin_override"
    override = ScoreEntry(
        event_id=original.event_id,
        pad_id=original.pad_id,
        run_id=original.run_id,
        competitor_id=original.competitor_id,
        judge_id=original.judge_id,
        device_id=session.device_id,
        source_type=source_type,
        status="submitted",
        submitted_at=now_utc(),
        original_score_entry_id=original.id,
        active_for_results=True,
        reason=payload.reason,
    )
    db.add(override)
    db.flush()
    total = 0.0
    for item_input in payload.items:
        criterion = db.get(Criterion, item_input.criterion_id)
        if not criterion:
            continue
        points = calculate_points(criterion, item_input.value)
        if points is not None:
            total += points
        db.add(
            ScoreItemEntry(
                score_entry_id=override.id,
                criterion_id=item_input.criterion_id,
                value=item_input.value,
                calculated_points=points,
                intentionally_blank=item_input.intentionally_blank,
            )
        )
    override.total = total
    competing_entries = db.scalars(
        select(ScoreEntry).where(
            ScoreEntry.event_id == session.event_id,
            ScoreEntry.run_id == original.run_id,
            ScoreEntry.judge_id == original.judge_id,
            ScoreEntry.status == "submitted",
            ScoreEntry.active_for_results.is_(True),
            ScoreEntry.id != override.id,
        )
    ).all()
    for entry in competing_entries:
        entry.active_for_results = False
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="score_override_created",
        entity_type="score_entry",
        entity_id=override.id,
        details={
            "original_score_entry_id": original.id,
            "run_id": run.id,
            "source_type": source_type,
            "total": total,
            "reason": payload.reason,
        },
    )
    db.commit()
    runs_by_id, judges_by_id = entry_context_maps(db, session.event_id, [override])
    return {"ok": True, "score": score_entry_payload(override, runs_by_id, judges_by_id)}


@router.post("/score-entries/{score_entry_id}/void")
def void_score_entry(score_entry_id: int, payload: ScoreVoidInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_owner(session)
    reason = (payload.reason or "").strip()
    if not reason:
        raise HTTPException(status_code=400, detail="A reason is required to delete judge scoring")
    entry = db.get(ScoreEntry, score_entry_id)
    if not entry or entry.event_id != session.event_id:
        raise HTTPException(status_code=404, detail="Score entry not found")
    old_status = entry.status
    old_active = bool(entry.active_for_results)
    entry.status = "voided"
    entry.active_for_results = False
    entry.reason = "\n".join(part for part in [entry.reason, f"Voided by owner: {reason}"] if part)
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        judge_id=entry.judge_id,
        device_id=session.device_id,
        action_type="score_entry_voided",
        entity_type="score_entry",
        entity_id=entry.id,
        details={
            "run_id": entry.run_id,
            "competitor_id": entry.competitor_id,
            "judge_id": entry.judge_id,
            "old_status": old_status,
            "old_active_for_results": old_active,
            "reason": reason,
        },
    )
    db.commit()
    return {"ok": True, "score_entry_id": entry.id, "status": entry.status}


def recovery_records_from_payload(raw_payload) -> list[dict]:
    if isinstance(raw_payload, list):
        return [item for item in raw_payload if isinstance(item, dict)]
    if isinstance(raw_payload, dict):
        if isinstance(raw_payload.get("records"), list):
            return [item for item in raw_payload["records"] if isinstance(item, dict)]
        if isinstance(raw_payload.get("score_records"), list):
            return [item for item in raw_payload["score_records"] if isinstance(item, dict)]
        if "run_id" in raw_payload and "items" in raw_payload:
            return [raw_payload]
    return []


def recovered_int(value) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def recovered_float(value) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


@router.post("/recovery-imports")
def create_recovery_import(payload: RecoveryImportInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_high_admin(session)
    recovery = RecoveryImport(
        event_id=session.event_id,
        imported_by_pin_account_id=session.pin_account_id,
        source_device_id=payload.source_device_id,
        raw_payload=payload.raw_payload,
        status="pending_review",
    )
    db.add(recovery)
    db.flush()
    imported_scores = 0
    for record in recovery_records_from_payload(payload.raw_payload):
        run_id = recovered_int(record.get("run_id"))
        items = record.get("items") or []
        if not run_id or not isinstance(items, list):
            continue
        run = db.get(Run, run_id)
        if not run or run.event_id != session.event_id:
            continue
        source_session_id = recovered_int(record.get("session_id"))
        source_session = db.get(JudgeSession, source_session_id) if source_session_id else None
        if source_session and source_session.event_id != session.event_id:
            source_session = None
        score = ScoreEntry(
            event_id=session.event_id,
            pad_id=run.pad_id,
            run_id=run.id,
            competitor_id=run.competitor_id,
            judge_id=source_session.judge_id if source_session else None,
            device_id=record.get("device_id") or payload.source_device_id,
            source_type="judge_device_recovery",
            status="submitted" if record.get("status") == "submitted" else "draft",
            submitted_at=now_utc() if record.get("status") == "submitted" else None,
            active_for_results=False,
            reason=f"Recovery import #{recovery.id}. Review before making official.",
        )
        db.add(score)
        db.flush()
        total = 0.0
        for raw_item in items:
            criterion_id = recovered_int(raw_item.get("criterion_id"))
            if not criterion_id:
                continue
            criterion = db.get(Criterion, criterion_id)
            if not criterion:
                continue
            value = recovered_float(raw_item.get("value"))
            points = calculate_points(criterion, value)
            if points is not None:
                total += points
            db.add(
                ScoreItemEntry(
                    score_entry_id=score.id,
                    criterion_id=criterion.id,
                    value=value,
                    calculated_points=points,
                    intentionally_blank=bool(raw_item.get("intentionally_blank")),
                )
            )
        score.total = total
        imported_scores += 1
    if imported_scores:
        recovery.status = "imported_pending_review"
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="recovery_import_created",
        entity_type="recovery_import",
        entity_id=recovery.id,
        details={"source_device_id": payload.source_device_id, "imported_scores": imported_scores},
    )
    db.commit()
    return {"ok": True, "recovery_import_id": recovery.id, "imported_scores": imported_scores, "status": recovery.status}


@router.get("/recovery-imports/recent")
def recent_recovery_scores(session_id: int, limit: int = 25, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_high_admin(session)
    entries = db.scalars(
        select(ScoreEntry)
        .options(selectinload(ScoreEntry.items).selectinload(ScoreItemEntry.criterion))
        .where(
            ScoreEntry.event_id == session.event_id,
            ScoreEntry.source_type == "judge_device_recovery",
        )
        .order_by(ScoreEntry.updated_at.desc(), ScoreEntry.id.desc())
        .limit(max(1, min(int(limit or 25), 50)))
    ).all()
    runs_by_id, judges_by_id = entry_context_maps(db, session.event_id, entries)
    return {"scores": [score_entry_payload(entry, runs_by_id, judges_by_id) for entry in entries]}


@router.post("/recovery-scores/{score_entry_id}/official")
def make_recovery_score_official(
    score_entry_id: int, payload: AdminStatusActionInput, db: Session = Depends(get_db)
) -> dict:
    session = get_session(db, payload.session_id)
    require_high_admin(session)
    entry = db.get(ScoreEntry, score_entry_id)
    if not entry or entry.event_id != session.event_id or entry.source_type != "judge_device_recovery":
        raise HTTPException(status_code=404, detail="Recovery score not found")
    if entry.status != "submitted":
        raise HTTPException(status_code=400, detail="Only submitted recovery scores can be made official")
    if not entry.judge_id:
        raise HTTPException(status_code=400, detail="Recovery score is not linked to a judge account")
    competing_entries = db.scalars(
        select(ScoreEntry).where(
            ScoreEntry.event_id == session.event_id,
            ScoreEntry.run_id == entry.run_id,
            ScoreEntry.judge_id == entry.judge_id,
            ScoreEntry.status == "submitted",
            ScoreEntry.active_for_results.is_(True),
            ScoreEntry.id != entry.id,
        )
    ).all()
    for competing in competing_entries:
        competing.active_for_results = False
    entry.active_for_results = True
    entry.reason = " ".join(
        part
        for part in [
            entry.reason or "",
            f"Made official from recovery review. {payload.note}".strip() if payload.note else "Made official from recovery review.",
        ]
        if part
    )
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="recovery_score_made_official",
        entity_type="score_entry",
        entity_id=entry.id,
        details={
            "run_id": entry.run_id,
            "judge_id": entry.judge_id,
            "deactivated_score_entry_ids": [item.id for item in competing_entries],
            "note": payload.note,
        },
    )
    db.commit()
    runs_by_id, judges_by_id = entry_context_maps(db, session.event_id, [entry])
    return {"ok": True, "score": score_entry_payload(entry, runs_by_id, judges_by_id)}


@router.get("/runs")
def list_admin_runs(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_admin(session)
    runs = db.scalars(select(Run).where(Run.event_id == session.event_id).order_by(Run.queue_position, Run.id)).all()
    score_counts = {
        (competitor_id, heat_number): int(count or 0)
        for competitor_id, heat_number, count in db.execute(
            select(Run.competitor_id, Run.heat_number, func.count(ScoreEntry.id))
            .join(ScoreEntry, ScoreEntry.run_id == Run.id)
            .where(
                Run.event_id == session.event_id,
                ScoreEntry.event_id == session.event_id,
                ScoreEntry.status.in_(["draft", "submitted"]),
            )
            .group_by(Run.competitor_id, Run.heat_number)
        ).all()
    }
    return {
        "runs": [
            {
                "id": run.id,
                "queue_position": run.queue_position,
                "run_number": run.run_number,
                "heat_number": run.heat_number,
                "run_type": run.run_type,
                "state": run.state,
                "include_in_results": run.include_in_results,
                "heat_score_count": score_counts.get((run.competitor_id, run.heat_number), 0),
                "competitor": {
                    "entry_number": run.competitor.entry_number,
                    "driver_name": run.competitor.driver_name,
                    "class_name": run.competitor.class_name,
                    "status": run.competitor.status,
                    "id": run.competitor.id,
                },
                "vehicle": {
                    "id": run.vehicle.id if run.vehicle else None,
                    "name": run.vehicle.name if run.vehicle else "",
                    "engine": run.vehicle.engine if run.vehicle else "",
                    "plate": run.vehicle.plate if run.vehicle else "",
                },
            }
            for run in runs
        ]
    }


@router.post("/run-order")
def update_run_order(payload: RunOrderInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_admin(session)
    changed = 0
    for item in payload.items:
        run = db.get(Run, item.run_id)
        if run and run.event_id == session.event_id:
            run.queue_position = item.queue_position
            changed += 1
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="run_order_updated",
        entity_type="run",
        details={"changed": changed},
    )
    db.commit()
    return {"ok": True, "changed": changed}


@router.get("/results")
def result_summary(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_admin(session)
    state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == session.event_id))
    settings = event_settings_record(db, session.event_id)
    runs = db.scalars(select(Run).where(Run.event_id == session.event_id).order_by(Run.queue_position, Run.id)).all()
    submitted_entries = db.scalars(
        select(ScoreEntry)
        .options(selectinload(ScoreEntry.items).selectinload(ScoreItemEntry.criterion))
        .where(
            ScoreEntry.event_id == session.event_id,
            ScoreEntry.status == "submitted",
            ScoreEntry.active_for_results.is_(True),
        )
        .order_by(ScoreEntry.updated_at.desc())
    ).all()
    runs_by_id = {run.id: run for run in runs}
    _, judges_by_id = entry_context_maps(db, session.event_id, submitted_entries)
    entries_by_run: dict[int, list[ScoreEntry]] = {}
    for entry in submitted_entries:
        entries_by_run.setdefault(entry.run_id, []).append(entry)

    rows = []
    for run in runs:
        entries = entries_by_run.get(run.id, [])
        total = sum(float(entry.total or 0) for entry in entries)
        adjusted_entries = [
            entry
            for entry in entries
            if entry.source_type in {"high_admin_override", "owner_override", "judge_device_recovery"}
            or entry.original_score_entry_id
        ]
        rows.append(
            {
                "run_id": run.id,
                "heat_number": run.heat_number,
                "competitor": f"#{run.competitor.entry_number} {run.competitor.driver_name}",
                "vehicle": run.vehicle.name if run.vehicle else "",
                "run_type": run.run_type,
                "state": run.state,
                "submitted_scores": len(entries),
                "total": total,
                "has_adjustments": bool(adjusted_entries),
                "adjustment_notes": [
                    entry.reason or f"{(entry.source_type or 'adjusted').replace('_', ' ')} entry #{entry.id}"
                    for entry in adjusted_entries
                ],
                "run": run_payload(run),
                "entries": [score_entry_payload(entry, runs_by_id, judges_by_id) for entry in entries],
                "criteria_totals": criteria_totals_payload(entries),
            }
        )
    current_heat = state.current_heat if state else 1
    existing_heats = sorted({row["heat_number"] or 1 for row in rows})
    next_heat = max(existing_heats + [current_heat]) + 1 if existing_heats or current_heat else 2
    heat_list = sorted(set(existing_heats + [current_heat, next_heat]))
    scoreboard = competitor_scoreboard(db, session.event_id, settings.result_mode)
    return {
        "results": rows,
        "current_heat": current_heat,
        "heats": heat_list,
        "result_mode": scoreboard["result_mode"],
        "result_mode_label": scoreboard["result_mode_label"],
        "overall_results": scoreboard["competitors"],
    }


@router.post("/current-heat/{heat_number}")
def set_current_heat(heat_number: int, session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_high_admin(session)
    if heat_number < 1:
        raise HTTPException(status_code=400, detail="Heat must be 1 or higher")
    state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == session.event_id))
    if not state:
        raise HTTPException(status_code=404, detail="Event state not found")
    old = state.current_heat
    state.current_heat = heat_number
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="current_heat_changed",
        entity_type="current_event_state",
        entity_id=state.id,
        details={"old_heat": old, "new_heat": heat_number},
    )
    db.commit()
    return {"ok": True, "current_heat": heat_number}


@router.get("/criteria")
def admin_criteria(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_admin(session)
    criteria_set = db.scalar(
        select(CriteriaSet).where(CriteriaSet.event_id == session.event_id, CriteriaSet.active.is_(True)).limit(1)
    )
    criteria = db.scalars(
        select(Criterion).where(Criterion.criteria_set_id == criteria_set.id).order_by(Criterion.display_order, Criterion.id)
    ).all()
    return {
        "criteria_set": {"id": criteria_set.id, "name": criteria_set.name},
        "criteria": [
            {
                "id": item.id,
                "name": item.name,
                "short_label": item.short_label,
                "type": item.type,
                "min_value": float(item.min_value) if item.min_value is not None else None,
                "max_value": float(item.max_value) if item.max_value is not None else None,
                "step_size": float(item.step_size),
                "points_per_unit": float(item.points_per_unit),
                "allows_negative": item.allows_negative,
                "required": item.required,
                "visible_to_judges": item.visible_to_judges,
                "visible_to_graphics": item.visible_to_graphics,
                "display_order": item.display_order,
                "active": item.active,
            }
            for item in criteria
        ],
    }


@router.post("/criteria")
def create_criterion(payload: CriterionInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_admin(session)
    criteria_set = db.scalar(
        select(CriteriaSet).where(CriteriaSet.event_id == session.event_id, CriteriaSet.active.is_(True)).limit(1)
    )
    criterion = Criterion(
        criteria_set_id=criteria_set.id,
        name=payload.name,
        short_label=payload.short_label,
        type=payload.type,
        min_value=payload.min_value,
        max_value=payload.max_value,
        step_size=payload.step_size,
        points_per_unit=payload.points_per_unit,
        allows_negative=payload.allows_negative,
        required=payload.required,
        visible_to_judges=payload.visible_to_judges,
        visible_to_graphics=payload.visible_to_graphics,
        display_order=payload.display_order,
        active=True,
    )
    db.add(criterion)
    db.flush()
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="criterion_created",
        entity_type="criterion",
        entity_id=criterion.id,
        details=payload.model_dump(exclude={"session_id"}),
    )
    db.commit()
    return {"ok": True, "criterion_id": criterion.id}


@router.put("/criteria/{criterion_id}")
def update_criterion(criterion_id: int, payload: CriterionInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_admin(session)
    criterion = db.get(Criterion, criterion_id)
    if not criterion:
        raise HTTPException(status_code=404, detail="Criterion not found")
    for field, value in payload.model_dump(exclude={"session_id"}).items():
        setattr(criterion, field, value)
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="criterion_updated",
        entity_type="criterion",
        entity_id=criterion.id,
        details=payload.model_dump(exclude={"session_id"}),
    )
    db.commit()
    return {"ok": True}


@router.delete("/criteria/{criterion_id}")
def delete_criterion(criterion_id: int, session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    criterion = db.get(Criterion, criterion_id)
    if not criterion:
        raise HTTPException(status_code=404, detail="Criterion not found")
    criterion.active = False
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="criterion_deleted",
        entity_type="criterion",
        entity_id=criterion.id,
    )
    db.commit()
    return {"ok": True}


@router.post("/criteria/{criterion_id}/toggle")
def toggle_criterion(criterion_id: int, session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_admin(session)
    criterion = db.get(Criterion, criterion_id)
    if not criterion:
        raise HTTPException(status_code=404, detail="Criterion not found")
    criterion.active = not criterion.active
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="criterion_toggled",
        entity_type="criterion",
        entity_id=criterion.id,
        details={"active": criterion.active},
    )
    db.commit()
    return {"ok": True, "active": criterion.active}


@router.get("/pin-accounts")
def list_pin_accounts(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    accounts = db.scalars(
        select(PinAccount).where(PinAccount.event_id == session.event_id).order_by(PinAccount.permission_level, PinAccount.role)
    ).all()
    return {
        "accounts": [
            {
                "id": account.id,
                "display_name": account.display_name,
                "pin": account.pin,
                "role": account.role,
                "permission_level": account.permission_level,
                "active": account.active,
            }
            for account in accounts
        ]
    }


@router.post("/pin-accounts")
def create_pin_account(payload: PinAccountInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_owner(session)
    validate_pin_payload(payload)

    existing = db.scalar(select(PinAccount).where(PinAccount.event_id == session.event_id, PinAccount.pin == payload.pin))
    if existing:
        raise HTTPException(status_code=400, detail="PIN already exists for this event")

    account = PinAccount(
        event_id=session.event_id,
        display_name=payload.display_name,
        pin=payload.pin,
        role=payload.role,
        permission_level=payload.permission_level,
        active=payload.active,
    )
    db.add(account)
    db.flush()

    if payload.role == "judge":
        db.add(Judge(event_id=session.event_id, pin_account_id=account.id, name=payload.display_name))

    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="pin_account_created",
        entity_type="pin_account",
        entity_id=account.id,
        details={"role": payload.role, "permission_level": payload.permission_level},
    )
    db.commit()
    return {"ok": True, "pin_account_id": account.id}


@router.put("/pin-accounts/{account_id}")
def update_pin_account(account_id: int, payload: PinAccountInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_owner(session)
    validate_pin_payload(payload)
    account = db.get(PinAccount, account_id)
    if not account or account.event_id != session.event_id:
        raise HTTPException(status_code=404, detail="PIN account not found")
    if account.role == "owner" and session.role != "owner":
        raise HTTPException(status_code=403, detail="Only owner can edit owner accounts")

    duplicate = db.scalar(
        select(PinAccount).where(
            PinAccount.event_id == session.event_id,
            PinAccount.pin == payload.pin,
            PinAccount.id != account.id,
        )
    )
    if duplicate:
        raise HTTPException(status_code=400, detail="PIN already exists for this event")

    if account.role == "owner" and (payload.role != "owner" or not payload.active):
        active_owners = db.query(PinAccount).filter(
            PinAccount.event_id == session.event_id,
            PinAccount.role == "owner",
            PinAccount.active.is_(True),
            PinAccount.id != account.id,
        ).count()
        if active_owners == 0:
            raise HTTPException(status_code=400, detail="Cannot remove the last active owner")

    old = {
        "display_name": account.display_name,
        "pin": account.pin,
        "role": account.role,
        "permission_level": account.permission_level,
        "active": account.active,
    }
    account.display_name = payload.display_name
    account.pin = payload.pin
    account.role = payload.role
    account.permission_level = payload.permission_level
    account.active = payload.active

    judge = db.scalar(select(Judge).where(Judge.pin_account_id == account.id))
    if payload.role == "judge":
        if judge:
            judge.name = payload.display_name
            judge.active = payload.active
            judge.status = "active" if payload.active else "removed"
        else:
            db.add(Judge(event_id=session.event_id, pin_account_id=account.id, name=payload.display_name, active=payload.active))
    elif judge:
        judge.active = False
        judge.status = "removed"

    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="pin_account_updated",
        entity_type="pin_account",
        entity_id=account.id,
        details={"old": old, "new": payload.model_dump(exclude={"session_id"})},
    )
    db.commit()
    return {"ok": True}


@router.delete("/pin-accounts/{account_id}")
def delete_pin_account(account_id: int, session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    account = db.get(PinAccount, account_id)
    if not account or account.event_id != session.event_id:
        raise HTTPException(status_code=404, detail="PIN account not found")
    if account.role == "owner" and session.role != "owner":
        raise HTTPException(status_code=403, detail="Only owner can delete owner accounts")
    if account.role == "owner":
        active_owners = db.query(PinAccount).filter(
            PinAccount.event_id == session.event_id,
            PinAccount.role == "owner",
            PinAccount.active.is_(True),
            PinAccount.id != account.id,
        ).count()
        if active_owners == 0:
            raise HTTPException(status_code=400, detail="Cannot delete the last active owner")

    account.active = False
    judge = db.scalar(select(Judge).where(Judge.pin_account_id == account.id))
    if judge:
        judge.active = False
        judge.status = "removed"
    active_sessions = db.scalars(select(JudgeSession).where(JudgeSession.pin_account_id == account.id)).all()
    for active_session in active_sessions:
        active_session.active = False
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="pin_account_deleted",
        entity_type="pin_account",
        entity_id=account.id,
        details={"soft_delete": True},
    )
    db.commit()
    return {"ok": True}


def normalize_header(value: str) -> str:
    return "".join(ch for ch in value.strip().lower() if ch.isalnum())


def pick(row: dict[str, str], *names: str) -> str:
    for name in names:
        value = row.get(normalize_header(name))
        if value not in {None, ""}:
            return str(value).strip()
    return ""


KNOWN_IMPORT_HEADERS = {
    normalize_header(value)
    for value in {
        "entry number",
        "entry",
        "number",
        "driver name",
        "driver",
        "name",
        "competitor",
        "car",
        "car name",
        "vehicle",
        "model",
        "engine",
        "engine type",
        "plate",
        "number plate",
        "registration",
        "class",
        "class label",
        "class code",
        "category",
        "sponsor",
        "notes",
        "run type",
        "raw source",
    }
}


def header_score(values: list[str]) -> int:
    return sum(1 for value in values if normalize_header(value) in KNOWN_IMPORT_HEADERS)


def table_rows_to_dicts(table_rows: list[list[str]]) -> list[dict[str, str]]:
    cleaned_rows = [[str(value or "").strip() for value in row] for row in table_rows]
    cleaned_rows = [row for row in cleaned_rows if any(value for value in row)]
    if not cleaned_rows:
        return []

    header_index = 0
    best_score = -1
    for index, row in enumerate(cleaned_rows[:10]):
        score = header_score(row)
        if score > best_score:
            best_score = score
            header_index = index
    if best_score <= 0:
        header_index = 0

    headers = []
    seen_headers = {}
    for index, value in enumerate(cleaned_rows[header_index]):
        header = normalize_header(value) or f"column{index + 1}"
        seen_headers[header] = seen_headers.get(header, 0) + 1
        if seen_headers[header] > 1:
            header = f"{header}{seen_headers[header]}"
        headers.append(header)

    parsed = []
    for source_row_number, values in enumerate(cleaned_rows[header_index + 1 :], start=header_index + 2):
        row = {headers[idx]: str(value or "").strip() for idx, value in enumerate(values) if idx < len(headers)}
        row["__source_row"] = str(source_row_number)
        parsed.append(row)
    return parsed


def infer_entry_number(row: dict[str, str], csv_row_number: int) -> tuple[str, str | None]:
    entry_number = pick(row, "entry number", "entry", "number", "car number", "competitor number")
    if entry_number:
        return entry_number, None

    driver_name = pick(row, "driver name", "driver", "name", "competitor")
    if driver_name and not driver_name.isdigit():
        inferred = str(csv_row_number - 1)
        return inferred, f"Row {csv_row_number}: inferred missing entry number as {inferred}"

    return "", None


def infer_driver_and_car(row: dict[str, str], entry_number: str, csv_row_number: int) -> tuple[str, str, str | None]:
    driver_name = pick(row, "driver name", "driver", "name", "competitor")
    car_name = pick(row, "car", "car name")
    if driver_name and not driver_name.isdigit():
        return driver_name, car_name, None

    raw_source = pick(row, "raw source")
    if not raw_source:
        return "", car_name, None

    tokens = raw_source.split()
    if tokens and tokens[0] == entry_number:
        tokens = tokens[1:]
    if len(tokens) < 2:
        return "", car_name, None

    inferred_driver = " ".join(tokens[:2])
    inferred_car = car_name
    if len(tokens) >= 3 and (not car_name or car_name.upper() == tokens[0].upper()):
        inferred_car = tokens[2]
    return inferred_driver, inferred_car, f"Row {csv_row_number}: inferred driver/car from raw source"


def parse_uploaded_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig")
        return table_rows_to_dicts(list(csv.reader(StringIO(text))))

    if filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="Excel import needs openpyxl. Rebuild the Docker image.") from exc
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        return table_rows_to_dicts([list(row) for row in rows])

    raise HTTPException(status_code=400, detail="Upload a CSV or XLSX file")


def build_import_preview(rows: list[dict[str, str]]) -> list[dict]:
    field_has_values = {
        "class": any(pick(row, "class label", "class", "category", "class code") for row in rows),
        "car_vehicle": any(pick(row, "car", "car name", "vehicle", "model") for row in rows),
        "engine": any(pick(row, "engine", "engine type") for row in rows),
        "plate": any(pick(row, "plate", "number plate", "registration") for row in rows),
        "sponsor": any(pick(row, "sponsor", "sponsors") for row in rows),
    }
    preview = []
    for index, row in enumerate(rows, start=2):
        source_row = int(row.get("__source_row") or index)
        entry_number, entry_note = infer_entry_number(row, source_row)
        driver_name, inferred_car_name, driver_note = infer_driver_and_car(row, entry_number, source_row)
        car_name = inferred_car_name or pick(row, "car", "car name")
        vehicle = pick(row, "vehicle", "model")
        issues = []
        warnings = []
        if entry_note:
            warnings.append(entry_note)
        if driver_note:
            warnings.append(driver_note)
        if not entry_number:
            issues.append("missing entry number")
        if not driver_name:
            issues.append("missing driver name")
        if field_has_values["car_vehicle"] and not car_name and not vehicle:
            warnings.append("missing car/vehicle")
        engine = pick(row, "engine", "engine type")
        if field_has_values["engine"] and not engine:
            warnings.append("missing engine")
        class_name = pick(row, "class label", "class", "category") or pick(row, "class code")
        if field_has_values["class"] and not class_name:
            warnings.append("missing class")
        if pick(row, "needs review").lower() in {"yes", "true", "1"}:
            review_note = pick(row, "review notes")
            warnings.append(f"source parser warning{': ' + review_note if review_note else ''}")
        notes = issues + warnings

        preview.append(
            {
                "import_row": source_row,
                "entry_number": entry_number,
                "driver_name": driver_name,
                "class_name": class_name,
                "car_name": car_name,
                "vehicle": vehicle,
                "engine": engine,
                "plate": pick(row, "plate", "number plate", "registration"),
                "sponsor": pick(row, "sponsor", "sponsors"),
                "notes": "; ".join(notes),
                "run_type": pick(row, "run type", "type") or "competition",
                "skip": False if entry_number or driver_name or car_name or vehicle else True,
                "needs_review": bool(issues),
                "has_warnings": bool(warnings),
                "critical_issues": issues,
                "warnings": warnings,
                "raw_source": pick(row, "raw source"),
            }
        )
    return preview


def annotate_import_preview(db: Session, event_id: int, preview: list[dict]) -> list[dict]:
    seen_entries: dict[str, int] = {}
    proposed_position = 0
    for row in preview:
        entry_number = (row.get("entry_number") or "").strip()
        if not row.get("skip"):
            proposed_position += 1
        row["queue_position"] = proposed_position if not row.get("skip") else None
        row["existing"] = None
        row["differences"] = []
        row["import_action"] = "review" if row.get("needs_review") else "new"

        if entry_number:
            seen_entries[entry_number] = seen_entries.get(entry_number, 0) + 1
            if seen_entries[entry_number] > 1:
                row["needs_review"] = True
                row.setdefault("critical_issues", []).append("duplicate entry number in import file")
                row["notes"] = "; ".join(filter(None, [row.get("notes"), "duplicate entry number in import file"]))
                row["import_action"] = "review"

        if row.get("skip"):
            row["import_action"] = "skip"
            continue
        if not entry_number or not (row.get("driver_name") or "").strip():
            row["import_action"] = "review"
            continue

        competitor = db.scalar(
            select(Competitor).where(Competitor.event_id == event_id, Competitor.entry_number == entry_number)
        )
        if not competitor:
            if row["import_action"] != "review":
                row["import_action"] = "new"
            continue

        run = db.scalar(
            select(Run)
            .where(Run.event_id == event_id, Run.competitor_id == competitor.id)
            .order_by(Run.run_number, Run.id)
            .limit(1)
        )
        vehicle = run.vehicle if run else None
        existing = {
            "competitor_id": competitor.id,
            "run_id": run.id if run else None,
            "entry_number": competitor.entry_number or "",
            "driver_name": competitor.driver_name or "",
            "class_name": competitor.class_name or "",
            "car_name": vehicle.name if vehicle else "",
            "engine": vehicle.engine if vehicle else "",
            "plate": vehicle.plate if vehicle else "",
            "run_type": run.run_type if run else "",
            "queue_position": run.queue_position if run else None,
        }
        row["existing"] = existing

        comparisons = [
            ("driver_name", "Driver"),
            ("class_name", "Class"),
            ("car_name", "Car"),
            ("engine", "Engine"),
            ("plate", "Plate"),
            ("run_type", "Run type"),
        ]
        differences = []
        for field, label in comparisons:
            incoming = str(row.get(field) or "").strip()
            current = str(existing.get(field) or "").strip()
            if incoming and incoming != current:
                differences.append({"field": field, "label": label, "from": current, "to": incoming})
        if run and row.get("queue_position") is not None and run.queue_position != row["queue_position"]:
            differences.append(
                {
                    "field": "queue_position",
                    "label": "Queue order",
                    "from": str(run.queue_position),
                    "to": str(row["queue_position"]),
                }
            )
        row["differences"] = differences
        if differences:
            row["import_action"] = "review" if row.get("needs_review") else "update"
        elif row["import_action"] != "review":
            row["import_action"] = "unchanged"
            row["skip"] = True
    return preview


@router.post("/competitor-import/preview")
async def preview_competitor_import(
    session_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> dict:
    session = get_session(db, session_id)
    require_admin(session)
    content = await file.read()
    rows = parse_uploaded_rows(file.filename or "", content)
    preview = annotate_import_preview(db, session.event_id, build_import_preview(rows))
    return {
        "ok": True,
        "filename": file.filename,
        "rows": preview,
        "needs_review": sum(1 for row in preview if row["needs_review"]),
        "warnings": sum(1 for row in preview if row.get("has_warnings") and not row["needs_review"]),
        "skipped_by_default": sum(1 for row in preview if row["skip"]),
        "new": sum(1 for row in preview if row["import_action"] == "new"),
        "updates": sum(1 for row in preview if row["import_action"] == "update"),
        "unchanged": sum(1 for row in preview if row["import_action"] == "unchanged"),
    }


@router.post("/competitor-import/commit")
def commit_competitor_import(payload: CompetitorImportCommit, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_admin(session)
    imported = 0
    updated = 0
    skipped = 0
    errors = []

    max_position = db.scalar(select(func.max(Run.queue_position)).where(Run.event_id == session.event_id)) or 0
    import_position = 0
    pad = get_event_pad(db, session.event_id)

    for row in payload.rows:
        if row.skip:
            skipped += 1
            continue
        entry_number = (row.entry_number or "").strip()
        driver_name = (row.driver_name or "").strip()
        if not entry_number or not driver_name:
            skipped += 1
            errors.append(f"Row {row.import_row}: missing entry number or driver name")
            continue
        import_position += 1

        competitor = db.scalar(
            select(Competitor).where(Competitor.event_id == session.event_id, Competitor.entry_number == entry_number)
        )
        if competitor:
            competitor.driver_name = driver_name
            competitor.class_name = row.class_name or competitor.class_name
            competitor.notes = row.notes or competitor.notes
            updated += 1
        else:
            competitor = Competitor(
                event_id=session.event_id,
                entry_number=entry_number,
                driver_name=driver_name,
                class_name=row.class_name or None,
                notes=row.notes or None,
                status="registered",
            )
            db.add(competitor)
            db.flush()
            imported += 1

        car_name = row.car_name or row.vehicle or "Vehicle TBC"
        engine = row.engine or None
        plate = row.plate or None
        sponsor = row.sponsor or None
        vehicle = db.scalar(
            select(Vehicle).where(
                Vehicle.event_id == session.event_id,
                Vehicle.name == car_name,
                Vehicle.plate == plate,
            )
        )
        if not vehicle:
            vehicle = Vehicle(
                event_id=session.event_id,
                name=car_name,
                engine=engine or None,
                plate=plate or None,
                sponsor=sponsor or None,
            )
            db.add(vehicle)
            db.flush()

        run = db.scalar(
            select(Run).where(
                Run.event_id == session.event_id,
                Run.competitor_id == competitor.id,
                Run.run_number == 1,
            )
        )
        position = import_position or max_position + 1
        if not run:
            max_position += 1
            run = Run(
                event_id=session.event_id,
                pad_id=pad.id,
                competitor_id=competitor.id,
                vehicle_id=vehicle.id,
                run_number=1,
                heat_number=max(1, int(row.heat_number or 1)),
                run_type=row.run_type or "competition",
                state=None,
                queue_position=position,
                include_in_results=(row.run_type or "competition") == "competition",
            )
            db.add(run)
            db.flush()
        else:
            run.vehicle_id = vehicle.id
            run.heat_number = max(1, int(row.heat_number or run.heat_number or 1))
            run.run_type = row.run_type or "competition"
            run.queue_position = position
            run.include_in_results = (row.run_type or "competition") == "competition"

        assignment = db.scalar(
            select(CompetitorVehicleAssignment).where(
                CompetitorVehicleAssignment.event_id == session.event_id,
                CompetitorVehicleAssignment.competitor_id == competitor.id,
                CompetitorVehicleAssignment.vehicle_id == vehicle.id,
                CompetitorVehicleAssignment.active.is_(True),
            )
        )
        if not assignment:
            db.add(
                CompetitorVehicleAssignment(
                    event_id=session.event_id,
                    competitor_id=competitor.id,
                    vehicle_id=vehicle.id,
                    run_id=run.id,
                    active=True,
                )
            )
        else:
            assignment.run_id = run.id

    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="competitor_imported",
        entity_type="competitor",
        details={"imported": imported, "updated": updated, "skipped": skipped, "errors": errors[:20]},
    )
    db.commit()
    return {"ok": True, "imported": imported, "updated": updated, "skipped": skipped, "errors": errors[:50]}


@router.post("/competitors")
def create_competitor(payload: CompetitorInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_admin(session)
    competitor = db.scalar(
        select(Competitor).where(Competitor.event_id == session.event_id, Competitor.entry_number == payload.entry_number)
    )
    if not competitor:
        competitor = Competitor(
            event_id=session.event_id,
            entry_number=payload.entry_number,
            driver_name=payload.driver_name,
            class_name=payload.class_name,
            notes=payload.notes,
            status="registered",
        )
        db.add(competitor)
        db.flush()
    else:
        competitor.driver_name = payload.driver_name
        competitor.class_name = payload.class_name
        competitor.notes = payload.notes

    vehicle = db.scalar(
        select(Vehicle).where(
            Vehicle.event_id == session.event_id,
            Vehicle.name == payload.car_name,
            Vehicle.plate == payload.plate,
        )
    )
    if not vehicle:
        vehicle = Vehicle(
            event_id=session.event_id,
            name=payload.car_name,
            engine=payload.engine,
            plate=payload.plate,
            sponsor=payload.sponsor,
        )
        db.add(vehicle)
        db.flush()

    position = payload.queue_position
    if position is None:
        position = (db.scalar(select(func.max(Run.queue_position)).where(Run.event_id == session.event_id)) or 0) + 1
    pad = get_event_pad(db, session.event_id)

    run = Run(
        event_id=session.event_id,
        pad_id=pad.id,
        competitor_id=competitor.id,
        vehicle_id=vehicle.id,
        run_number=1,
        heat_number=max(1, int(payload.heat_number or 1)),
        run_type=payload.run_type,
        state=payload.run_state,
        queue_position=position,
        include_in_results=payload.run_type == "competition",
    )
    db.add(run)
    db.flush()
    db.add(
        CompetitorVehicleAssignment(
            event_id=session.event_id,
            competitor_id=competitor.id,
            vehicle_id=vehicle.id,
            run_id=run.id,
            active=True,
        )
    )
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="competitor_created",
        entity_type="competitor",
        entity_id=competitor.id,
        details={"run_id": run.id},
    )
    db.commit()
    return {"ok": True, "competitor_id": competitor.id, "run_id": run.id}


@router.put("/runs/{run_id}")
def update_run_details(run_id: int, payload: CompetitorInput, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, payload.session_id)
    require_admin(session)
    run = db.get(Run, run_id)
    if not run or run.event_id != session.event_id:
        raise HTTPException(status_code=404, detail="Run not found")

    competitor = run.competitor
    competitor.entry_number = payload.entry_number
    competitor.driver_name = payload.driver_name
    competitor.class_name = payload.class_name
    competitor.notes = payload.notes

    vehicle = run.vehicle
    if not vehicle:
        vehicle = Vehicle(event_id=session.event_id, name=payload.car_name or "Vehicle TBC")
        db.add(vehicle)
        db.flush()
        run.vehicle_id = vehicle.id
    vehicle.name = payload.car_name or "Vehicle TBC"
    vehicle.engine = payload.engine
    vehicle.plate = payload.plate
    vehicle.sponsor = payload.sponsor

    run.run_type = payload.run_type
    run.heat_number = max(1, int(payload.heat_number or 1))
    run.state = payload.run_state
    if payload.queue_position is not None:
        run.queue_position = payload.queue_position
    run.include_in_results = payload.run_type == "competition"

    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="run_competitor_updated",
        entity_type="run",
        entity_id=run.id,
        details={"competitor_id": competitor.id, "vehicle_id": vehicle.id},
    )
    db.commit()
    return {"ok": True}


@router.post("/run/{run_id}/type/{run_type}")
def update_current_run_type(run_id: int, run_type: str, session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_admin(session)
    run = db.get(Run, run_id)
    if not run or run.event_id != session.event_id:
        raise HTTPException(status_code=404, detail="Run not found")
    if run_type not in RUN_TYPES:
        raise HTTPException(status_code=400, detail="Invalid run type")

    old_type = run.run_type
    run.run_type = run_type
    run.include_in_results = run_type == "competition"
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="run_type_changed",
        entity_type="run",
        entity_id=run.id,
        details={"old": old_type, "new": run_type, "include_in_results": run.include_in_results},
    )
    db.commit()
    return {"ok": True, "run_id": run.id, "run_type": run.run_type, "include_in_results": run.include_in_results}


@router.get("/backup")
def backup_data(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    return build_backup_payload(db, reason="manual_download")


@router.get("/backup/health")
def backup_health(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    return backup_health_snapshot()


@router.post("/backup/write")
def write_server_backup(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    path = write_backup(db, reason="manual")
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="server_backup_written",
        entity_type="backup",
        details={"path": str(path)},
    )
    db.commit()
    return {"ok": True, "path": str(path), "health": backup_health_snapshot()}


@router.post("/restore")
async def restore_backup(
    session_id: int,
    confirmation: str = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    if confirmation != "RESTORE KAIRIX":
        raise HTTPException(status_code=400, detail="Confirmation must be RESTORE KAIRIX")

    content = await file.read()
    try:
        payload = json.loads(content.decode("utf-8"))
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail="Backup file must be JSON") from exc

    tables = payload.get("tables", {})
    if not isinstance(tables, dict):
        raise HTTPException(status_code=400, detail="Backup file is missing tables")

    for model in reversed(BACKUP_MODELS):
        db.query(model).delete()
    db.flush()

    inserted = {}
    for model in BACKUP_MODELS:
        table_name = model.__tablename__
        rows = tables.get(table_name, [])
        if not rows:
            inserted[table_name] = 0
            continue
        cleaned_rows = [coerce_row_for_model(model, row) for row in rows]
        db.execute(model.__table__.insert(), cleaned_rows)
        inserted[table_name] = len(cleaned_rows)

    db.commit()
    return {
        "ok": True,
        "message": "Backup restored. You may need to log in again if sessions were replaced.",
        "inserted": inserted,
    }


def coerce_row_for_model(model, row: dict) -> dict:
    output = {}
    columns = {col.name: col for col in model.__table__.columns}
    for key, value in row.items():
        if key not in columns:
            continue
        column = columns[key]
        if value is None:
            output[key] = None
            continue
        if "DATETIME" in str(column.type).upper() and isinstance(value, str):
            try:
                output[key] = datetime.fromisoformat(value.replace("Z", "+00:00"))
            except ValueError:
                output[key] = value
        else:
            output[key] = value
    return output


@router.get("/db/summary")
def database_summary(session_id: int, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    return {
        "tables": [
            {"table": name, "count": db.query(model).count()}
            for name, model in sorted(BACKUP_TABLES.items())
        ]
    }


@router.get("/db/table/{table_name}")
def database_table(table_name: str, session_id: int, limit: int = 100, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    model = BACKUP_TABLES.get(table_name)
    if not model:
        raise HTTPException(status_code=404, detail="Unknown table")
    limit = max(1, min(limit, 500))
    rows = []
    for obj in db.scalars(select(model).limit(limit)).all():
        row = {}
        for col in model.__table__.columns:
            value = getattr(obj, col.name)
            row[col.name] = value.isoformat() if hasattr(value, "isoformat") else value
        rows.append(row)
    return {"table": table_name, "limit": limit, "rows": rows}


@router.post("/db/run-state-cleanup")
def cleanup_run_states(session_id: int, mode: str = "scoring_to_registered", db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    state = db.scalar(select(CurrentEventState).where(CurrentEventState.event_id == session.event_id))
    current_id = state.official_current_run_id if state else None
    runs = db.scalars(select(Run).where(Run.event_id == session.event_id)).all()
    changed = 0
    for run in runs:
        if run.id == current_id:
            continue
        if mode == "scoring_to_registered" and run.state == "scoring":
            run.state = None
            changed += 1
        elif mode == "all_non_current_to_registered" and run.state not in {"finished", "skipped", "withdrawn"}:
            run.state = None
            changed += 1
    log_action(
        db,
        event_id=session.event_id,
        actor_pin_account_id=session.pin_account_id,
        device_id=session.device_id,
        action_type="run_state_cleanup",
        entity_type="run",
        details={"mode": mode, "changed": changed},
    )
    db.commit()
    return {"ok": True, "changed": changed}


@router.post("/danger/clear-database")
def clear_database(session_id: int, confirmation: str, db: Session = Depends(get_db)) -> dict:
    session = get_session(db, session_id)
    require_owner(session)
    if confirmation != "CLEAR KAIRIX":
        raise HTTPException(status_code=400, detail="Confirmation must be CLEAR KAIRIX")

    for model in reversed(BACKUP_MODELS):
        db.query(model).delete()
    db.commit()
    seed_initial_data(db, include_sample_competitors=False)
    return {"ok": True}
